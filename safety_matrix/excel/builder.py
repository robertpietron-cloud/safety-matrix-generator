# =============================================================================
# Project : Safety Matrix Generator (GuardLogix SafetyTask -> Cause & Effect)
# File    : safety_matrix/excel/builder.py
# Author  : Robert Pietroń
# Revision: 1.1.0 (2026-10-21)
# Changes :
#   1.0.0 (2026-10-18) - First universal release (layout of PortalSaw matrix
#                        rev 1.4 generalised): Start (language switch), Group
#                        Interlocks, MOTOR, VALVE, Legend, Safety functions
#                        (reaction time, PL fields), Test record, Devices,
#                        Outputs, C&E paths, Validation report, Dictionary.
#   1.1.0 (2026-10-21) - Channels copied from physical inputs shown as
#                        'copy ← input'; modules/discrepancy test use the
#                        physical channel; constant-contact bypass (code N) and
#                        scan-order lag in the paths, comments and test record;
#                        scan lag included in the reaction time; ROUT/CROUT
#                        with a start / seal-in on its own rung = manual restart.
# =============================================================================
"""Excel workbook generator."""
from __future__ import annotations

import re
from collections import OrderedDict
from typing import Dict, List, Optional

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

from .. import __version__
from .. import instructions as INS
from ..analysis import MATRIX_CODES, Analysis, Effect
from ..audit import ALARM, DEAD, VERIFY, Finding
from ..config import Config, Layout
from ..i18n import EXCEL_LANG, LANG_NAMES, ML, Tr
from ..ladder import seal_in_legs
from ..model import io_parts
from .base import (AREA_FILL, BOX, CENTER, CODE_STYLE, HDR_FILL, INFILL, LEFT, LEFT_TOP, META_FILL,
                   RESULT_STYLE, TYPE_FILL, VERT, ZONE_FILL, Book, F, IFont, inject_cached_values)

TRIP = ('X', 'X1', 'XB', 'XM')
DRIVE_KINDS = ('STO', 'SS1', 'SBC', 'SLS')
SOURCES = [
    ('Rockwell 1756-RM012 – Connection Reaction Time Limit',
     'https://www.rockwellautomation.com/en-au/docs/technical/logix5000/_online/1756-rm012/guardlogix-5580-and-compact-guardlogix-5580-safety/reaction-times/connection-reaction-time-limit.html'),
    ('Rockwell 1756-UM900 – Factors That Affect Logix Reaction-time Components',
     'https://www.rockwellautomation.com/en-ie/docs/technical/logix5000/_online/1756-um900/controllogix-5590-controller-user-manual-ditamap/reaction-times/factors-that-affect-logix-reaction-time-components.html'),
    ('Rockwell – Safety instructions (Studio 5000 Logix Designer)',
     'https://www.rockwellautomation.com/en-au/docs/studio-5000-logix-designer/37-00/contents-ditamap/instruction-set/safety-instructions.html'),
]


class MatrixWorkbook:
    def __init__(self, A: Analysis, findings: List[Finding], layout: Layout, cfg: Config, tr: Tr):
        self.A, self.F, self.Ly, self.cfg, self.tr = A, findings, layout, cfg, tr
        self.B = Book(tr, cfg.default_language, cfg.author)
        self.T = tr
        self.p = A.p
        self.multi = self.B.multi
        self.SN = {k: self.B.sheet_name('sheet.' + k) for k in ('start', 'legend', 'sf', 'test', 'dev', 'out', 'path', 'rep', 'dict')}
        self.effects = [A.effect(k) for k in layout.ordered_effects()]
        self.safe = [e for e in self.effects if not e.info]
        self.devs = layout.ordered_devices()
        self._prepare()

    # ================================================================= text helpers
    def t(self, key, **kw) -> ML:
        return self.tr(key, **kw)

    def name(self) -> ML:
        if self.cfg.title:
            return self.tr.ml(self.cfg.title)
        return self.t('doc.title', name=self.p.controller or self.p.source_name)

    def title_block(self) -> ML:
        tb = self.t('doc.block', author=self.cfg.author or '—', rev=self.cfg.revision, date=self.cfg.date)
        return tb + self.t('doc.block_lang') if self.multi else tb

    def source(self) -> ML:
        return self.t('doc.source', src=self.p.source_name or '—', proc=self.p.processor or '—',
                      ver=self.p.software_version or '—', period=_g(self.A.task_period))

    def loc(self, l) -> str:
        return self.A.loc_label(l)

    def t_nconst(self, tag: str, effs: List[Effect]) -> ML:
        """Test expectation for outputs on which the device is bridged by a constant contact (code N)."""
        bits = sorted({b for e in effs for b in self.A.cells[(tag, e.key)].bypass})
        return self.t('t.exp.nconst', outs=self.compress(effs), bits=', '.join(bits))

    def compress(self, effs: List[Effect]) -> ML:
        keys = {e.key for e in effs}
        def build(l):
            parts = []
            for g in self.Ly.groups:
                h = [self.Ly.effect_label(k) for k in g.effects if k in keys]
                if not h:
                    continue
                gn = self.B.txt(g.name, l)
                if len(h) == len(g.effects) and len(h) > 2:
                    parts.append(f"{gn}: {self.tr.s(l, 'w.all', n=len(h))}")
                else:
                    parts.append(f"{gn}: " + ', '.join(h))
            return '; '.join(parts)
        return self.tr.of(build)

    def src_text(self, op: str) -> ML:
        A = self.A
        k = A.source_kind(op)
        if k == 'std':
            return self.tr.ml(f'{op} ← {A.std_trace(op) or A.std_source(op) or "std"}')
        if k == 'input':
            c = A.io_comment(op)
            return self.tr.ml(f"{op} '{c}'" if c else op)
        if k in ('const0', 'const1'):
            return self.t('src.const', op=op, v=1 if k == 'const1' else 0)
        if k == 'never':
            return self.t('src.never', op=op)
        d = A.describe(op)
        return self.tr.ml(f"{op} '{' '.join(d.split())}'" if d else op)

    def reset_text(self, tag: str):
        """(sources ML, writer rungs str) of a reset operand."""
        if not tag:
            return self.tr.ml('—'), ''
        if INS.looks_literal(tag):
            return self.t('src.literal', v=tag), ''
        srcs, locs = self.A.reset_sources(tag)
        items = [self.src_text(s) for s in srcs]
        if not items:
            items = [self.src_text(tag)]
        return self.tr.join('; ', items), ', '.join(self.loc(l) for l in locs)

    # ================================================================= prepared data
    def _prepare(self):
        A, Ly = self.A, self.Ly
        self.cells = A.cells
        self.EM: Dict[str, dict] = {}
        for e in self.effects:
            g = Ly.group_of(e.key)
            m = dict(key=e.key, label=Ly.effect_label(e.key), group=g.name if g else self.t('group.other'),
                     desc=Ly.eff_desc.get(e.key, self.tr.ml(e.key)), kind=e.kind, valve=e.key in Ly.valves)
            rung = A.effect_rung(e)
            m['out_rung'] = self.loc(rung)
            en, en_loc, en_expr = A.enable_of(e)
            m['enable'], m['enable_rung'], m['en_expr'] = en, self.loc(en_loc), en_expr
            m['kind_ml'] = self.t('kindshort.' + e.kind)
            if e.kind in DRIVE_KINDS:
                mod = A.drive_module(e)
                m['addr'] = f"{mod.name} ({mod.address})" if mod and mod.address else (mod.name if mod else e.key)
            elif e.info:
                m['addr'] = e.key
            else:
                mod = A.module_of(e.outputs[0])
                ad = self.p.top_adapter(mod) if mod else None
                outs = _compress_outputs(e.outputs)
                m['addr'] = outs + (f' ({ad.address})' if ad and ad.address else '')
            cells = [self.cells[(d, e.key)] for d in self.devs if (d, e.key) in self.cells and self.cells[(d, e.key)].code != 'N']
            m['cat'] = '—' if e.info else A.stop_category(cells)
            if e.instr:
                oi = A.lm.outputs[e.instr]
                spec = INS.SAFETY_INSTRUCTIONS[oi.instr]
                fb1, fb2 = oi.roles.get('Feedback1', ''), oi.roles.get('Feedback2', '')
                fbt = oi.roles.get('FeedbackType', '')
                m['kind_ml'] = self.t('kindshort.' + e.kind) + (f' ({fbt})' if fbt and not fbt.isdigit() else '')
                cm = A.io_comment(fb1)
                fbs = fb1 + (f' / {fb2}' if fb2 and fb2 != fb1 else '')
                m['edm'] = self.tr.ml(f"{fbs} ({fbt}) – '{cm}'") if cm else self.tr.ml(f'{fbs} ({fbt}) – ') + self.t('w.nocomment')
                m['fb'], m['fbtype'], m['fbcomment'] = fb1, fbt, cm
                rst = oi.roles.get(spec.reset, '')
                srcs, rr = self.reset_text(rst)
                m['reset_tag'], m['reset_rr'] = rst, rr
                m['reset'] = self.tr.ml(f'{rst} [{rr}] ') + self.t('rs.rout_fault')
                r_oi = A.rung(oi.loc)
                legs = seal_in_legs(r_oi[2], e.instr + '.O1') if r_oi else None
                start = [op for leg in (legs or []) if not any(x.startswith(e.instr + '.') for x in leg)
                         for op in leg if not A.is_status(op)]
                # start / seal-in on the ROUT rung itself: restart needs that start or reset signal
                m['restart'] = self.t('restart.rout_seal', tags=', '.join(start)) if start else self.t('restart.rout')
                m['rout'] = oi
            elif e.info:
                m['edm'], m['reset'], m['restart'] = self.tr.ml('—'), self.tr.ml('—'), self.t('restart.follow')
            else:
                manual, auto = A.restart_legs(e)
                m['edm'] = self.t('edm.cip') if e.kind in DRIVE_KINDS else self.tr.ml('—')
                parts = []
                for t in manual:
                    _, rr = self.reset_text(t)
                    parts.append(f'{t} [{rr}]' if rr else t)
                so_reset = ''
                if e.kind in DRIVE_KINDS:
                    drv = e.key.split(':')[0]
                    w = A.lm.writes.get(f'{drv}:SO.Reset')
                    if w:
                        so_reset = self.loc(w[0][0])
                rs = self.tr.ml(' ∥ '.join(parts)) if parts else self.t('rs.none')
                if so_reset:
                    rs = rs + self.t('rs.so_reset', rung=so_reset)
                m['reset'] = rs
                m['manual'], m['auto'] = manual, auto
                if auto:
                    m['restart'] = self.t('restart.auto', tags=', '.join(auto))
                elif manual:
                    m['restart'] = self.t('restart.manual')
                else:
                    m['restart'] = self.t('restart.follow')
            mo, con = A.output_connection(e)
            m['crtl'] = con.reaction_time_limit_ms if (con and not e.info) else None
            self.EM[e.key] = m

        self.DM: Dict[str, dict] = {}
        for z in Ly.zones:
            for tag in z.devices:
                d = A.devices[tag]
                m = dict(tag=tag, zone=z.name, zno=z.number, id=d.id, type=d.type, desc=Ly.dev_desc[tag], instr=d.instr,
                         cha=d.ch_a + (f' ← {d.src_a}' if d.src_a else ''), chb=d.ch_b + (f' ← {d.src_b}' if d.src_b else ''),
                         routine=d.loc.routine if d.loc else '—', rung=d.loc.rung if d.loc else '—', kind=d.kind)
                m['reset_type'] = {True: 'MANUAL', False: 'AUTOMATIC'}.get(d.reset_manual, '—')
                m['reset_tag'] = d.reset or '—'
                m['reset_src'], m['reset_rr'] = self.reset_text(d.reset) if d.reset else (self.tr.ml('—'), '')
                if d.mute and not INS.looks_literal(d.mute):
                    std = sorted(f for f in A.ev.cone(d.mute)[1] if A.leaf_class(f) == 'std')
                    m['mute'] = d.mute + (' ← ' + ', '.join(std) if std else '')
                else:
                    m['mute'] = ''
                m['n_eff'] = d.n_eff
                cells = [self.cells[(tag, e.key)] for e in self.safe if (tag, e.key) in self.cells and self.cells[(tag, e.key)].code in TRIP]
                m['cat'] = A.stop_category(cells)
                st = []
                if d.n_eff == 0:
                    st.append('st.dead')
                if d.reset_manual is False:
                    st.append('st.auto')
                if d.equal_channels and d.kind == 'instr':
                    st.append('st.equal')
                if d.const_channels:
                    st.append('st.const')
                if d.std_channels:
                    st.append('st.std')
                if m['mute'] and '←' in m['mute']:
                    st.append('st.mute')
                if d.kind == 'direct':
                    st.append('st.direct')
                if d.kind == 'external':
                    st.append('st.external')
                if d.heuristic and d.kind == 'instr':
                    st.append('st.heuristic')
                m['status_keys'] = st
                m['status'] = self.tr.join('; ', [self.t(k) for k in st]) if st else self.tr.ml('OK')
                self.DM[tag] = m

    def code(self, dev, eff) -> Optional[str]:
        return self.A.code(dev, eff)

    def hits(self, dev, codes=TRIP) -> List[Effect]:
        return [e for e in self.safe if self.code(dev, e.key) in codes]

    # ================================================================= matrix helpers
    def path_text(self, c) -> ML:
        s = ' → '.join(f'{t} [{self.loc(l)}]' for t, l in c.path)
        out = self.tr.ml(s)
        if c.bypass:
            out = out + '\n' + self.t('w.constbypass' if c.raw == 'N' else 'w.bypass', bits=', '.join(c.bypass))
        if c.lag:
            out = out + '\n' + self.t('w.lag', n=c.lag, ms=_g((self.A.task_period or 0) * c.lag))
        if c.conditions:
            out = out + '\n' + self.t('w.conditions', bits=', '.join(c.conditions))
        if c.partners:
            out = out + '\n' + self.t('w.partners', devs=', '.join(c.partners))
        if c.delay_ms is not None:
            out = out + '\n' + self.t('w.delay', ms=_g(c.delay_ms))
        return out

    def put_code(self, ws, r, col, dev, eff, comments=True):
        code = self.code(dev, eff)
        cell = self.B.put(ws, r, col, code, align=CENTER)
        if code:
            bg, fg, _ = CODE_STYLE[code]
            cell.fill = PatternFill('solid', fgColor=bg)
            cell.font = F(bold=True, color=fg)
            if comments and self.cfg.comments:
                self.B.comment(cell, self.path_text(self.cells[(dev, eff)]))
        return code

    def present_codes(self) -> List[str]:
        used = {self.code(d, e.key) for d in self.devs for e in self.effects} - {None}
        return [c for c in MATRIX_CODES if c in used]

    def corner_legend(self, ws, rng, read_key):
        parts = [(self.name() + '\n', IFont(11, b=True)), (self.title_block() + '\n', IFont(8, b=True)),
                 (self.source() + '\n\n', IFont(8, i=True)), (self.t('lg.title') + '\n', IFont(9, b=True))]
        for code in self.present_codes() + ['']:
            parts.append((code.ljust(3) if code else self.t('code.blank'), IFont(9, b=True, color=CODE_STYLE[code][2])))
            parts.append((self.tr.ml('  – ') + self.t(f'code.{code or "none"}.short') + '\n', IFont(8)))
        types = [t for t in INS.DEVICE_TYPES if any(self.DM[d]['type'] == t for d in self.devs)]
        parts.append((self.t('lg.ids'), IFont(8, b=True)))
        parts.append((self.tr.of(lambda l: ' · '.join(f"{INS.TYPE_ABBR[t]} = {self.tr.s(l, 'type.' + t).lower()}" for t in types))
                      + self.t('lg.id_suffix') + '\n', IFont(8)))
        parts.append((self.t('lg.abbr_head'), IFont(8, b=True)))
        parts.append((self.t('lg.abbr_line') + '\n\n', IFont(8)))
        parts.append((self.t(read_key) + '\n', IFont(8, i=True)))
        parts.append((self.t('lg.full_note'), IFont(8, b=True)))
        self.B.rich(ws, rng, parts)

    def legend_block(self, ws, row, code_col, text_from, text_to, read_key, text_width):
        B = self.B

        def line_h(txt):
            t = B.txt(txt)
            return max(15, 12 * (len(t) // max(int(text_width * 1.15), 20) + 1))

        def header(title):
            nonlocal row
            B.put(ws, row, code_col, title, F(bold=True, color='FFFFFF'), HDR_FILL, LEFT)
            ws.merge_cells(start_row=row, start_column=code_col, end_row=row, end_column=text_to)
            ws.row_dimensions[row].height = 16
            row += 1

        def item(code, txt, fill=None, font=None):
            nonlocal row
            B.put(ws, row, code_col, code, font or F(bold=True), fill or TYPE_FILL, CENTER)
            B.put(ws, row, text_from, txt, F(size=8), None, LEFT)
            for c in range(text_from + 1, text_to + 1):
                ws.cell(row=row, column=c).border = BOX
            ws.merge_cells(start_row=row, start_column=text_from, end_row=row, end_column=text_to)
            ws.row_dimensions[row].height = line_h(txt)
            row += 1

        header(self.t('lg.title') + '   (' + self.title_block() + ')')
        item(self.t('lg.read'), self.t(read_key), font=F(bold=True, size=8))
        header(self.t('lg.codes'))
        for code in self.present_codes() + ['']:
            bg, fg, _ = CODE_STYLE[code]
            item(code or self.t('code.blank'), self.t(f'code.{code or "none"}.full'), PatternFill('solid', fgColor=bg), F(bold=True, color=fg))
        header(self.t('lg.ids_head'))
        for t in INS.DEVICE_TYPES:
            if any(self.DM[d]['type'] == t for d in self.devs):
                item(INS.TYPE_ABBR[t], self.t('type.' + t) + ' – ' + self.t('typedesc.' + t))
        item(self.t('lg.format'), self.t('lg.id_format'), font=F(bold=True, size=8))
        header(self.t('lg.kinds'))
        for k in OrderedDict.fromkeys(e.kind for e in self.effects):
            item(self.t('kindshort.' + k), self.t('kind.' + k), font=F(bold=True, size=7))
        header(self.t('lg.abbr_main'))
        for i, (k, v) in enumerate(self.tr.table('abbr')):
            if i in (0, 1, 2, 4, 5, 6, 7, 8):
                item(k, v, font=F(bold=True, size=7))
        return row

    def add_code_rules(self, ws, sqref):
        codes = ','.join(MATRIX_CODES)
        dv = DataValidation(type='list', formula1=f'"{codes}"', allow_blank=True, showErrorMessage=True,
                            errorTitle=self.B.txt(self.t('dv.err_title'))[:32], error=self.B.txt(self.t('dv.err', codes=codes))[:255],
                            showInputMessage=True, promptTitle=self.B.txt(self.t('dv.prompt_title'))[:32],
                            prompt=self.B.txt(self.t('dv.prompt'))[:255])
        for rng_ in sqref.split():
            dv.add(rng_)
        ws.add_data_validation(dv)
        for code in MATRIX_CODES:
            bg, fg, _ = CODE_STYLE[code]
            ws.conditional_formatting.add(sqref, CellIsRule(operator='equal', formula=[f'"{code}"'],
                                          fill=PatternFill(start_color=bg, end_color=bg, fill_type='solid'),
                                          font=Font(name='Arial', bold=True, color=fg)))
        ws.conditional_formatting.add(sqref, CellIsRule(operator='equal', formula=['""'],
                                      fill=PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')))

    # ================================================================= build
    def build(self, path: str):
        wb = self.B.wb
        self.ws_start = wb.active
        self.ws_start.title = self.SN['start']
        self.sheet_gi()
        self.sheet_matrix('MOTOR', [e for e in self.effects if not self.EM[e.key]['valve']], 'read.motor')
        valves = [e for e in self.effects if self.EM[e.key]['valve']]
        if valves:
            self.sheet_matrix('VALVE', valves, 'read.motor')
        else:
            self.sheet_valve_empty()
        self.sheet_legend()
        self.sheet_sf()
        self.sheet_tests()
        self.sheet_devices()
        self.sheet_outputs()
        self.sheet_paths()
        self.sheet_report()
        self.sheet_start()
        if self.multi:
            self.B.write_dictionary(self.SN['dict'], self.SN['start'])
        self.finish(path)

    # ----------------------------------------------------------------- Group Interlocks
    def sheet_gi(self):
        B = self.B
        ws = B.wb.create_sheet('Group Interlocks')
        self.ws_gi = ws
        self.corner_legend(ws, 'A1:C2', 'read.gi')
        B.put(ws, 3, 1, self.t('gi.effect'), F(bold=True), TYPE_FILL, CENTER)
        B.put(ws, 3, 2, self.t('gi.tag'), F(bold=True), TYPE_FILL, CENTER)
        B.put(ws, 3, 3, self.t('gi.enable'), F(bold=True), TYPE_FILL, CENTER)
        col = 4
        dev_col = OrderedDict()
        for zi, z in enumerate(self.Ly.zones):
            if zi:
                col += 2
            B.put(ws, 1, col, z.name, F(bold=True, color='FFFFFF'), HDR_FILL, VERT)
            ws.merge_cells(start_row=1, start_column=col, end_row=3, end_column=col)
            ws.column_dimensions[CL(col)].width = 5
            col += 2
            prev = None
            for tag in z.devices:
                dm = self.DM[tag]
                if prev and dm['type'] != prev:
                    col += 1
                prev = dm['type']
                B.put(ws, 1, col, dm['type'], F(bold=True, size=8), TYPE_FILL, VERT)
                B.put(ws, 2, col, dm['desc'], F(size=8), None, VERT)
                B.put(ws, 3, col, f"{dm['id']}\n{tag}", F(bold=True, size=8), ZONE_FILL, VERT)
                dev_col[tag] = col
                ws.column_dimensions[CL(col)].width = 4.6
                col += 1
        last_col = max(col - 1, 4)
        ws.row_dimensions[1].height = 95
        ws.row_dimensions[2].height = 230
        ws.row_dimensions[3].height = 120
        for c, w in zip('ABC', (46, 24, 40)):
            ws.column_dimensions[c].width = w
        for c in range(4, last_col + 1):
            if ws.column_dimensions[CL(c)].width in (None, 13):
                ws.column_dimensions[CL(c)].width = 1.8
        row = 4
        eff_rows = []
        for g in self.Ly.groups:
            B.put(ws, row, 1, g.name, F(bold=True), AREA_FILL)
            for c in range(2, last_col + 1):
                ws.cell(row=row, column=c).fill = AREA_FILL
            row += 1
            for k in g.effects:
                em = self.EM[k]
                B.put(ws, row, 1, em['desc'], align=LEFT)
                B.put(ws, row, 2, em['label'], F(bold=True), align=LEFT)
                B.put(ws, row, 3, f"{em['enable']}  [{em['enable_rung']}]", F(size=8), align=LEFT)
                for tag, c in dev_col.items():
                    self.put_code(ws, row, c, tag, k, comments=False)
                ws.row_dimensions[row].height = max(24, 11 * (len(em['enable']) // 48 + 1))
                eff_rows.append(row)
                row += 1
        meta = [('gi.reset_row', lambda d: {'MANUAL': 'M', 'AUTOMATIC': 'A'}.get(d['reset_type'], '—')),
                ('gi.cat_row', lambda d: d['cat']),
                ('gi.count_row', lambda d: d['n_eff'])]
        for key, fn in meta:
            B.put(ws, row, 1, self.t(key), F(bold=True, size=8), META_FILL, LEFT)
            B.put(ws, row, 2, None, fill=META_FILL)
            B.put(ws, row, 3, None, fill=META_FILL)
            for tag, c in dev_col.items():
                v = fn(self.DM[tag])
                cell = B.put(ws, row, c, v, F(size=8, bold=True), META_FILL, CENTER)
                if v in ('A', 0):
                    cell.font = F(size=8, bold=True, color='C00000')
            row += 1
        row += 1
        if not any(self.EM[e.key]['valve'] for e in self.effects):
            B.put(ws, row, 1, self.t('gi.valve_head'), F(bold=True), AREA_FILL)
            row += 1
            B.put(ws, row, 1, self.t('gi.valve_none'), F(italic=True), align=LEFT, border=False)
            row += 2
        self.legend_block(ws, row, 2, 3, last_col, 'read.gi', text_width=200)
        if eff_rows and dev_col:
            self.add_code_rules(ws, _ranges(eff_rows, list(dev_col.values())))
        ws.freeze_panes = 'D4'
        ws.print_title_rows = '1:3'
        self.gi_rows = eff_rows
        self.gi_cols = list(dev_col.values())

    # ----------------------------------------------------------------- MOTOR / VALVE
    def sheet_matrix(self, title, effs: List[Effect], read_key):
        B = self.B
        wm = B.wb.create_sheet(title)
        self.corner_legend(wm, 'A1:C1', read_key)
        B.put(wm, 2, 1, self.t('m.type'), F(bold=True), TYPE_FILL, CENTER)
        B.put(wm, 2, 2, self.t('m.desc'), F(bold=True), TYPE_FILL, CENTER)
        B.put(wm, 2, 3, self.t('m.id'), F(bold=True), TYPE_FILL, CENTER)
        keys = {e.key for e in effs}
        col = 4
        eff_col = OrderedDict()
        gi = 0
        for g in self.Ly.groups:
            ks = [k for k in g.effects if k in keys]
            if not ks:
                continue
            if gi:
                col += 3
            gi += 1
            B.put(wm, 1, col, g.name, F(bold=True, color='FFFFFF'), HDR_FILL, VERT)
            wm.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
            wm.column_dimensions[CL(col)].width = 5
            col += 1
            for k in ks:
                em = self.EM[k]
                B.put(wm, 1, col, em['desc'], F(size=8), None, VERT)
                B.put(wm, 2, col, em['label'], F(bold=True, size=8), ZONE_FILL, VERT)
                eff_col[k] = col
                wm.column_dimensions[CL(col)].width = 4.6
                col += 1
        m_last = max(col - 1, 4)
        for c in range(4, m_last + 1):
            if wm.column_dimensions[CL(c)].width in (None, 13):
                wm.column_dimensions[CL(c)].width = 1.8
        wm.row_dimensions[1].height = 230
        wm.row_dimensions[2].height = 120
        for c, w in zip('ABC', (17, 62, 33)):
            wm.column_dimensions[c].width = w
        row = 3
        first = row
        dev_rows = []
        for z in self.Ly.zones:
            B.put(wm, row, 1, z.name, F(bold=True), ZONE_FILL)
            for c in range(2, m_last + 1):
                wm.cell(row=row, column=c).fill = ZONE_FILL
            row += 1
            for tag in z.devices:
                dm = self.DM[tag]
                B.put(wm, row, 1, dm['type'], F(size=8), align=LEFT)
                B.put(wm, row, 2, dm['desc'], F(size=8), align=LEFT)
                B.put(wm, row, 3, f"{dm['id']}  {tag}", F(bold=True, size=8), align=LEFT)
                for k, c in eff_col.items():
                    self.put_code(wm, row, c, tag, k)
                wm.row_dimensions[row].height = 26
                dev_rows.append(row)
                row += 1
        last = row - 1
        row += 1
        B.put(wm, row, 1, self.t('m.meta_head'), F(bold=True), META_FILL)
        row += 1
        for key, k in [('m.kind', 'kind_ml'), ('m.addr', 'addr'), ('m.rung', 'out_rung'), ('m.enable', 'enable'),
                       ('m.edm', 'edm'), ('m.reset', 'reset'), ('m.restart', 'restart'), ('m.cat', 'cat')]:
            B.put(wm, row, 2, self.t(key), F(bold=True, size=8), META_FILL, LEFT)
            for ek, c in eff_col.items():
                B.put(wm, row, c, self.EM[ek][k], F(size=7), META_FILL, VERT)
            wm.row_dimensions[row].height = 150 if k in ('enable', 'edm', 'reset', 'restart') else 70
            row += 1
        row += 1
        B.put(wm, row, 2, self.t('m.count'), F(bold=True, size=8), META_FILL, LEFT)
        for ek, c in eff_col.items():
            rng = f'{CL(c)}{first}:{CL(c)}{last}'
            f = '=' + '+'.join(f'COUNTIF({rng},"{x}")' for x in ('X', 'X1', 'XB', 'XM', 'S', 'SB', 'SM'))
            n = sum(1 for d in self.devs if self.code(d, ek) in ('X', 'X1', 'XB', 'XM', 'S', 'SB', 'SM'))
            B.put_number_formula(wm, row, c, f, n, font=F(size=8, bold=True), fill=META_FILL, align=CENTER)
        row += 3
        self.legend_block(wm, row, 1, 2, 3, read_key, text_width=95)
        if dev_rows and eff_col:
            self.add_code_rules(wm, _ranges(dev_rows, list(eff_col.values())))
        wm.freeze_panes = 'D3'
        wm.print_title_rows = '1:2'
        if title == 'MOTOR':
            self.motor_range = (first, last, m_last)

    def sheet_valve_empty(self):
        B = self.B
        wv = B.wb.create_sheet('VALVE')
        B.put(wv, 1, 1, self.name(), F(bold=True, size=11), border=False)
        B.put(wv, 1, 2, self.title_block() + '\n' + self.t('v.legend_ref'), F(size=8), None, LEFT, border=False)
        wv.row_dimensions[1].height = 30
        B.put(wv, 1, 4, self.t('v.none_head'), F(bold=True, color='FFFFFF'), HDR_FILL, LEFT)
        wv.merge_cells('D1:J1')
        B.put(wv, 2, 1, self.t('m.type'), F(bold=True), TYPE_FILL, CENTER)
        B.put(wv, 2, 2, self.t('m.desc'), F(bold=True), TYPE_FILL, CENTER)
        B.put(wv, 2, 3, self.t('m.id'), F(bold=True), TYPE_FILL, CENTER)
        outs = [self.EM[e.key]['label'] for e in self.safe if e.kind in ('ROUT', 'CROUT', 'OUT')]
        B.put(wv, 2, 4, self.t('v.none_text', outs=', '.join(outs) or '—'), F(italic=True, size=8), None, LEFT, border=False)
        wv.merge_cells('D2:J2')
        wv.row_dimensions[2].height = 40
        row = 3
        for z in self.Ly.zones:
            B.put(wv, row, 1, z.name, F(bold=True), ZONE_FILL)
            row += 1
            for tag in z.devices:
                dm = self.DM[tag]
                B.put(wv, row, 1, dm['type'], F(size=8))
                B.put(wv, row, 2, dm['desc'], F(size=8), align=LEFT)
                B.put(wv, row, 3, f"{dm['id']}  {tag}", F(bold=True, size=8))
                row += 1
        for c, w in zip('ABC', (17, 62, 26)):
            wv.column_dimensions[c].width = w

    # ----------------------------------------------------------------- Devices
    def sheet_devices(self):
        B = self.B
        wd = B.wb.create_sheet(self.SN['dev'])
        hdr = ['d.zone', 'd.id', 'd.type', 'd.tag', 'd.desc', 'd.instr', 'd.cha', 'd.chb', 'd.routine', 'd.rung',
               'd.reset_type', 'd.reset_tag', 'd.reset_src', 'd.reset_rung', 'd.mute', 'd.n_eff', 'd.cat', 'd.status']
        for i, h in enumerate(hdr, 1):
            B.put(wd, 1, i, self.t(h), F(bold=True, color='FFFFFF'), HDR_FILL, CENTER)
        r = 2
        for tag in self.devs:
            d = self.DM[tag]
            vals = [self.t('w.zone_n', n=d['zno']), d['id'], d['type'], tag, d['desc'], d['instr'], d['cha'], d['chb'],
                    d['routine'], d['rung'], d['reset_type'], d['reset_tag'], d['reset_src'], d['reset_rr'], d['mute'],
                    d['n_eff'], d['cat'], d['status']]
            for i, v in enumerate(vals, 1):
                cell = B.put(wd, r, i, v, F(size=8), None, LEFT_TOP)
                if i == len(vals) and d['status_keys']:
                    cell.font = F(size=8, bold=True, color='C00000')
            r += 1
        for i, w in enumerate([10, 9, 15, 26, 48, 9, 22, 22, 10, 6, 11, 26, 60, 18, 30, 10, 9, 34], 1):
            wd.column_dimensions[CL(i)].width = w
        wd.freeze_panes = 'E2'
        wd.auto_filter.ref = f'A1:{CL(len(hdr))}{max(2, r - 1)}'

    # ----------------------------------------------------------------- Outputs
    def sheet_outputs(self):
        B = self.B
        we = B.wb.create_sheet(self.SN['out'])
        hdr = ['o.group', 'o.tag', 'o.desc', 'o.kind', 'o.addr', 'o.rung', 'o.enable', 'o.enable_rung', 'o.logic', 'o.edm',
               'o.reset', 'o.restart', 'o.cat', 'o.crtl']
        for i, h in enumerate(hdr, 1):
            B.put(we, 1, i, self.t(h), F(bold=True, color='FFFFFF'), HDR_FILL, CENTER)
        r = 2
        for e in self.effects:
            m = self.EM[e.key]
            vals = [m['group'], m['label'], m['desc'], m['kind_ml'], m['addr'], m['out_rung'], m['enable'], m['enable_rung'],
                    m['en_expr'], m['edm'], m['reset'], m['restart'], m['cat'], m['crtl'] if m['crtl'] is not None else '—']
            for i, v in enumerate(vals, 1):
                B.put(we, r, i, v, F(size=8), None, LEFT_TOP)
            r += 1
        aux = [e for e in self.A.effects if not e.in_matrix]
        if aux:
            r += 1
            B.put(we, r, 1, self.t('o.aux_head'), F(bold=True, color='FFFFFF'), HDR_FILL, LEFT)
            we.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
            r += 1
            for e in aux:
                B.put(we, r, 2, e.key, F(size=8, bold=True), None, LEFT_TOP)
                B.put(we, r, 3, self.t('o.reason.' + (e.reason or 'no_device')), F(size=8), None, LEFT_TOP)
                B.put(we, r, 6, self.loc(self.A.effect_rung(e)), F(size=8), None, LEFT_TOP)
                B.put(we, r, 9, self.A.permissive_text(e.var), F(size=8), None, LEFT_TOP)
                r += 1
        for i, w in enumerate([16, 24, 44, 16, 24, 14, 36, 14, 70, 44, 60, 40, 12, 12], 1):
            we.column_dimensions[CL(i)].width = w
        we.freeze_panes = 'C2'

    # ----------------------------------------------------------------- Paths
    def sheet_paths(self):
        B = self.B
        wp = B.wb.create_sheet(self.SN['path'])
        hdr = ['p.dev', 'p.eff', 'p.code', 'p.path', 'p.extra']
        for i, h in enumerate(hdr, 1):
            B.put(wp, 1, i, self.t(h), F(bold=True, color='FFFFFF'), HDR_FILL, CENTER)
        r = 2
        for e in self.effects:
            for d in self.devs:
                c = self.cells.get((d, e.key))
                if not c:
                    continue
                B.put(wp, r, 1, d, F(size=8), None, LEFT_TOP)
                B.put(wp, r, 2, self.EM[e.key]['label'], F(size=8), None, LEFT_TOP)
                bg, fg, _ = CODE_STYLE[c.code]
                B.put(wp, r, 3, c.code, F(size=8, bold=True, color=fg), PatternFill('solid', fgColor=bg), CENTER)
                B.put(wp, r, 4, ' → '.join(f'{t} [{self.loc(l)}]' for t, l in c.path), F(size=8), None, LEFT_TOP)
                if c.raw == 'N':
                    extra = '; '.join(f'{b} = 1 (const)' for b in c.bypass)
                else:
                    extra = '; '.join(f'{b} = {self.A.std_trace(b) or "std"}' for b in c.bypass) or ', '.join(c.partners) \
                        or ', '.join(c.conditions)
                if c.delay_ms is not None:
                    extra = (extra + '; ' if extra else '') + f'TOF {_g(c.delay_ms)} ms'
                if c.lag:
                    extra = (extra + '; ' if extra else '') + f'+{c.lag} × {_g(self.A.task_period)} ms (scan order)'
                B.put(wp, r, 5, extra, F(size=8), None, LEFT_TOP)
                r += 1
        for i, w in enumerate([26, 26, 6, 110, 60], 1):
            wp.column_dimensions[CL(i)].width = w
        wp.freeze_panes = 'A2'
        wp.auto_filter.ref = f'A1:E{max(2, r - 1)}'

    # ----------------------------------------------------------------- Report
    def finding_text(self, f: Finding) -> ML:
        return self.t('f.' + f.key, **f.params)

    def finding_rec(self, f: Finding) -> ML:
        if f.severity == DEAD:
            return self.t('cls.' + (f.cls or 'dead'))
        key = 'f.' + f.key + '.rec'
        return self.t(key, **f.params) if self.tr.has(key) else self.tr.ml('')

    def sheet_report(self):
        B = self.B
        wr = B.wb.create_sheet(self.SN['rep'])
        B.put(wr, 1, 1, self.t('r.title'), F(bold=True, size=12), border=False)
        B.put(wr, 2, 1, self.title_block() + '   •   ' + self.source(), F(size=8, bold=True), border=False)
        order = ' → '.join(r for _, r in self.A.lm.executed)
        B.put(wr, 3, 1, self.t('r.note', order=order), F(size=8, italic=True), border=False)
        r = 5
        for sev, title, cols, color in ((ALARM, 'r.sec_alarm', 'r.cols_alarm', 'C00000'), (DEAD, 'r.sec_dead', 'r.cols_dead', '7F7F7F'),
                                        (VERIFY, 'r.sec_verify', 'r.cols_verify', 'BF8F00')):
            items = [f for f in self.F if f.severity == sev]
            B.put(wr, r, 1, self.t(title, n=len(items)), F(bold=True, color='FFFFFF'), PatternFill('solid', fgColor=color), LEFT)
            wr.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            r += 1
            for i, h in enumerate(self.tr.table(cols)[0], 1):
                B.put(wr, r, i, h, F(bold=True), TYPE_FILL, CENTER)
            r += 1
            if not items:
                B.put(wr, r, 2, self.t('r.none'), F(size=8, italic=True), None, LEFT_TOP)
                r += 1
            for n, f in enumerate(items, 1):
                txt = self.finding_text(f)
                if f.priority:
                    txt = self.t('r.priority') + txt
                vals = [n, f.routine, f.rungs, f.element, txt, self.finding_rec(f)]
                for i, v in enumerate(vals, 1):
                    B.put(wr, r, i, v, F(size=8), None, CENTER if i == 1 else LEFT_TOP)
                L = max(len(B.txt(txt)), len(B.txt(vals[5])) * 95 // 60)
                wr.row_dimensions[r].height = min(400, max(24, 11 * (L // 105 + 1)))
                r += 1
            r += 1
        for i, w in enumerate([5, 16, 20, 40, 95, 60], 1):
            wr.column_dimensions[CL(i)].width = w
        self.n_alarm = sum(1 for f in self.F if f.severity == ALARM)
        self.n_verify = sum(1 for f in self.F if f.severity == VERIFY)
        self.n_dead = sum(1 for f in self.F if f.severity == DEAD)

    # ----------------------------------------------------------------- Legend
    def sheet_legend(self):
        B = self.B
        wl = B.wb.create_sheet(self.SN['legend'])
        self.ws_leg = wl
        B.put(wl, 1, 1, self.t('l.title'), F(bold=True, size=12), border=False)
        B.put(wl, 2, 1, self.title_block() + '   •   ' + self.source(), F(size=8, bold=True), border=False)
        r = 4

        def lsec(title, cols):
            nonlocal r
            B.put(wl, r, 1, title, F(bold=True, color='FFFFFF'), HDR_FILL, LEFT)
            wl.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            r += 1
            for i, h in enumerate(cols, 1):
                if h is not None:
                    B.put(wl, r, i, h, F(bold=True, size=8), TYPE_FILL, CENTER)
            r += 1

        def lrow(vals, fill=None, font=None, h=None):
            nonlocal r
            for i, v in enumerate(vals, 1):
                if v is None:
                    continue
                B.put(wl, r, i, v, (font if i == 1 and font else F(size=8, bold=(i == 1))), fill if i == 1 else None,
                      CENTER if i == 1 else LEFT)
            longest = max((len(B.txt(v)) for v in vals if v is not None and not (isinstance(v, str) and v.startswith('='))), default=10)
            wl.row_dimensions[r].height = h or max(16, 12 * (longest // 105 + 1))
            r += 1

        mf, ml_, mlast = getattr(self, 'motor_range', (3, 3, 4))
        rng = f"MOTOR!$D${mf}:${CL(mlast)}${ml_}"
        gi_rng = f"'Group Interlocks'!$D${min(self.gi_rows or [5])}:${CL(max(self.gi_cols or [4]))}${max(self.gi_rows or [5])}"
        hdr = self.tr.table('l.cols_codes')[0]
        lsec(self.t('l.s1'), hdr)
        motor_keys = [e.key for e in self.effects if not self.EM[e.key]['valve']]
        for code in MATRIX_CODES + ['']:
            bg, fg, _ = CODE_STYLE[code]
            row = [code or self.t('code.blank'), self.t(f'code.{code or "none"}.short'), self.t(f'code.{code or "none"}.full'), None, None]
            lrow(row, PatternFill('solid', fgColor=bg), F(bold=True, color=fg), h=30)
            if code:
                n_m = sum(1 for d in self.devs for k in motor_keys if self.code(d, k) == code)
                n_g = sum(1 for d in self.devs for e in self.effects if self.code(d, e.key) == code)
                B.put_number_formula(wl, r - 1, 4, f'=COUNTIF({rng},"{code}")', n_m, font=F(size=8), align=CENTER)
                B.put_number_formula(wl, r - 1, 5, f'=COUNTIF({gi_rng},"{code}")', n_g, font=F(size=8), align=CENTER)
        r += 1
        lsec(self.t('l.s2'), [self.t('l.sheet'), None, self.t('l.axes'), None, None])
        lrow(['Group Interlocks', None, self.t('read.gi')])
        lrow(['MOTOR', None, self.t('read.motor')])
        lrow(['VALVE', None, self.t('read.valve')])
        r += 1
        lsec(self.t('l.s3'), [self.t('l.id_abbr'), self.t('l.type_hdr'), self.t('l.meaning'), self.t('l.instr'), self.t('l.count')])
        for t in INS.DEVICE_TYPES:
            n = sum(1 for d in self.devs if self.DM[d]['type'] == t)
            if not n:
                continue
            instrs = ', '.join(sorted({self.DM[d]['instr'] for d in self.devs if self.DM[d]['type'] == t}))
            lrow([INS.TYPE_ABBR[t], t, self.t('typedesc.' + t), instrs, n], TYPE_FILL)
        lrow([self.t('lg.format'), None, self.t('lg.id_format')], TYPE_FILL)
        r += 1
        lsec(self.t('l.s4'), [self.t('l.zone_no'), self.t('l.id_prefix'), self.t('l.zone_name'), self.t('l.count'), None])
        for z in self.Ly.zones:
            lrow([self.t('w.zone_n', n=z.number), f'xx {z.number}.n', z.name, len(z.devices)], ZONE_FILL)
        r += 1
        lsec(self.t('l.s5'), [self.t('l.kind'), None, self.t('l.meaning'), self.t('l.count'), None])
        for k in OrderedDict.fromkeys(e.kind for e in self.effects):
            lrow([self.t('kindshort.' + k), None, self.t('kind.' + k), sum(1 for e in self.effects if e.kind == k)], TYPE_FILL, F(bold=True, size=8))
        r += 1
        lsec(self.t('l.s6'), [self.t('l.instr'), None, self.t('l.meaning'), None, None])
        used = list(OrderedDict.fromkeys([d.instr for d in self.A.lm.devices.values()] + [o.instr for o in self.A.lm.outputs.values()]
                                         + [m.instr for m in self.A.lm.mutings.values()] + [t.kind for t in self.A.lm.timers.values()]))
        for k in used + ['OSF / OSR', 'XIC / XIO', 'OTE / OTL / OTU', '.O1', '.FP']:
            key = 'ins.' + k
            lrow([k, None, self.t(key) if self.tr.has(key) else self.t('ins.unknown')], TYPE_FILL, F(bold=True, size=8))
        r += 1
        lsec(self.t('l.s7'), [self.t('l.mark'), None, self.t('l.meaning'), None, None])
        for k, v in self.tr.table('abbr'):
            lrow([k, None, v], TYPE_FILL, F(bold=True, size=8))
        r += 1
        lsec(self.t('l.s8'), [self.t('l.sample'), self.t('l.element'), self.t('l.meaning'), None, None])
        for (bg, fg), (name, txt) in zip([('1F3864', 'FFFFFF'), ('D9E1F2', '000000'), ('FCE4D6', '000000'), ('EDEDED', '000000'),
                                          ('F2F2F2', 'C00000')], self.tr.table('colors')):
            lrow(['Abc', name, txt], PatternFill('solid', fgColor=bg), F(bold=True, color=fg))
        r += 1
        B.put(wl, r, 1, self.t('l.s9'), F(bold=True, color='FFFFFF'), HDR_FILL, LEFT)
        wl.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        r += 1
        notes = [self.t('l.n.method'), self.t('l.n.cat1') if any(t.kind in INS.OFF_DELAY_TIMERS for t in self.A.lm.timers.values())
                 else self.t('l.n.cat0'), self.t('l.n.edm'), self.t('l.n.rout'), self.t('l.n.comment'), self.t('l.n.ids'),
                 self.t('l.n.tool', version=__version__)]
        for t in notes:
            B.put(wl, r, 1, '•', border=False, align=CENTER)
            B.put(wl, r, 2, t, F(size=8), align=LEFT, border=False)
            wl.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            wl.row_dimensions[r].height = 36
            r += 1
        r += 1
        lsec(self.t('l.s10'), [self.t('l.sheet'), None, self.t('l.contents'), None, None])
        for k, txt in self.sheet_list():
            lrow([k, None, txt], TYPE_FILL, F(bold=True, size=8))
        r += 1
        lsec(self.t('l.s11'), [self.t('l.mark'), None, self.t('l.meaning'), None, None])
        lrow([self.t('l.yellow'), None, self.t('input_note')], INFILL, F(bold=True, size=8))
        for code in ('OK', 'NOK', 'N/A'):
            bg, fg = RESULT_STYLE[code]
            lrow([code, None, self.t('result.' + code)], PatternFill('solid', fgColor=bg), F(bold=True, color=fg))
        lrow([self.t('l.dropdown'), None, self.t('l.dropdown_txt')], TYPE_FILL, F(bold=True, size=7))
        r += 1
        lsec(self.t('l.s12'), [self.t('l.rev'), self.t('l.date'), self.t('l.changes'), self.t('l.author'), None])
        for rv in self.revisions():
            lrow([rv[0], rv[1], rv[3], rv[2]], TYPE_FILL, F(bold=True, size=8))
        r += 1
        lsec(self.t('l.s13'), [self.t('l.no'), None, self.t('l.doc'), None, None])
        for i, (t, url) in enumerate(SOURCES, 1):
            lrow([str(i), None, f'{t}\n{url}'], TYPE_FILL, F(bold=True, size=8), h=28)
        for i, w in enumerate([18, 30, 105, 14, 16], 1):
            wl.column_dimensions[CL(i)].width = w
        wl.freeze_panes = 'A4'

    def revisions(self):
        out = []
        for rv in self.cfg.revisions:
            out.append((str(rv.get('rev', '')), str(rv.get('date', '')), str(rv.get('author', self.cfg.author or '')),
                        self.tr.ml(rv.get('text', ''))))
        if not any(r[0] == self.cfg.revision for r in out):
            out.append((self.cfg.revision, self.cfg.date, self.cfg.author or '—', self.t('l.rev_auto', version=__version__)))
        return out

    def sheet_list(self):
        names = [(self.SN['start'], 'sl.start'), ('Group Interlocks', 'sl.gi'), ('MOTOR', 'sl.motor'), ('VALVE', 'sl.valve'),
                 (self.SN['legend'], 'sl.legend'), (self.SN['sf'], 'sl.sf'), (self.SN['test'], 'sl.test'), (self.SN['dev'], 'sl.dev'),
                 (self.SN['out'], 'sl.out'), (self.SN['path'], 'sl.path'), (self.SN['rep'], 'sl.rep')]
        if self.multi:
            names.append((self.SN['dict'], 'sl.dict'))
        return [(n, self.t(k)) for n, k in names]

    # ----------------------------------------------------------------- Safety functions
    def sheet_sf(self):
        B, A = self.B, self.A
        wf = B.wb.create_sheet(self.SN['sf'])
        B.put(wf, 1, 1, self.t('sf.title'), F(bold=True, size=12), border=False)
        B.put(wf, 2, 1, self.title_block() + '   •   ' + self.source(), F(size=8, bold=True), border=False)
        B.put(wf, 3, 1, self.t('sf.method', period=_g(A.task_period), watchdog=_g(A.task_watchdog)), F(size=8), None, LEFT, border=False)
        wf.merge_cells('A3:S3')
        wf.row_dimensions[3].height = 36
        B.put(wf, 4, 1, self.t('input_note') + ' ' + self.t('sf.input_hint'), F(size=8, italic=True), INFILL, LEFT, border=False)
        wf.merge_cells('A4:S4')
        wf.row_dimensions[4].height = 26
        HR = 6
        for i, h in enumerate(self.tr.table('sf.cols')[0], 1):
            B.put(wf, HR, i, h, F(bold=True, color='FFFFFF', size=8), HDR_FILL, CENTER)
        wf.row_dimensions[HR].height = 42
        per, wd = A.task_period or 0, A.task_watchdog or 0
        fr = HR + 1
        for n, tag in enumerate(self.devs, 1):
            d, dev = self.DM[tag], A.devices[tag]
            hx = self.hits(tag)
            hxc = self.hits(tag, ('XC',))
            hx_t = hx + [e for e in hxc if e not in hx]
            codes = {self.code(tag, e.key) for e in self.safe} - {None}
            if dev.const_channels:
                struct = self.t('sf.s.const')
            elif dev.kind == 'direct':
                struct = self.t('sf.s.direct')
            elif dev.kind == 'external':
                struct = self.t('sf.s.external')
            elif dev.equal_channels:
                struct = self.t('sf.s.single', instr=dev.instr)
            elif dev.std_channels:
                struct = self.t('sf.s.std', instr=dev.instr)
            else:
                struct = self.t('sf.s.two', instr=dev.instr)
            if d['mute'] and '←' in d['mute']:
                struct = struct + self.t('sf.s.mute')
            if 'XB' in codes:
                struct = struct + self.t('sf.s.xb')
            if 'XC' in codes:
                struct = struct + self.t('sf.s.xc')
            if 'XM' in codes:
                struct = struct + self.t('sf.s.xm')
            mod, con = A.input_connection(dev)
            out_max = max((c.reaction_time_limit_ms or 0 for e in hx_t for c in [A.output_connection(e)[1]] if c), default=None)
            notes = []
            if hxc:
                partners = sorted({p for e in hxc for p in self.cells[(tag, e.key)].partners})
                notes.append(self.t('sf.n.xc', effs=', '.join(self.EM[e.key]['label'] for e in hxc), partners=', '.join(partners)))
            if not hx_t:
                notes.append(self.t('sf.n.nonfunc') if 'N' in codes else self.t('sf.n.none'))
            if con and (con.reaction_time_limit_ms or 0) >= self.cfg.crtl_priority_ms:
                notes.append(self.t('sf.n.crtl', ms=_g(self.cfg.crtl_priority_ms)))
            if dev.std_channels:
                notes.append(self.t('sf.n.std'))
            delays = sorted({self.cells[(tag, e.key)].delay_ms for e in hx_t if self.cells[(tag, e.key)].delay_ms is not None})
            if delays:
                notes.append(self.t('sf.n.delay', ms='/'.join(_g(x) for x in delays)))
            lag = max((self.cells[(tag, e.key)].lag for e in hx_t), default=0)
            if lag:
                notes.append(self.t('sf.n.lag', n=lag, ms=_g(per * lag)))
            outs = self.compress(hx)
            if hxc:
                outs = outs + (self.tr.ml('; ') if hx else self.tr.ml('')) + self.t('sf.only_xc', effs=', '.join(self.EM[e.key]['label'] for e in hxc))
            if not hx and not hxc:
                outs = self.tr.ml('—')
            short = d['desc'].map(lambda s: s.split('  [')[0])
            modtxt = (f"{mod.tag_prefix} {mod.catalog}: RPI {_g(con.rpi_ms)} ms / TM {_g(con.timeout_multiplier)} / "
                      f"NDM {_g(con.network_delay_multiplier)} %") if (mod and con) else 'n/a'
            vals = [f'SF-{n:02d}', d['id'], tag, self.t('fn.' + d['type']) + ': ' + short, outs, d['cat'] if hx_t else '—', struct, modtxt,
                    con.reaction_time_limit_ms if (con and hx_t) else None, (dev.filter_ms or 0) if hx_t else None,
                    per * (1 + lag) + wd if hx_t else None, out_max, None, None, None, None, None, None,
                    self.tr.join(' ', notes) if notes else '']
            for i, v in enumerate(vals, 1):
                cell = B.put(wf, fr, i, v, F(size=8, bold=(i in (1, 3))), None, LEFT_TOP if i not in (6, 9, 10, 11, 12, 15) else CENTER)
                if i in (13, 14, 16, 17, 18):
                    cell.fill = INFILL
                if i in (9, 10, 11, 12, 13, 14, 15):
                    cell.number_format = '0.0'
                if i == 7 and '[ALARM' in B.txt(v):
                    cell.font = F(size=8, bold=True, color='C00000')
            if hx_t and con:
                tot = sum(x for x in (con.reaction_time_limit_ms, dev.filter_ms or 0, per + wd, out_max) if x)
                B.put_number_formula(wf, fr, 15, f'=IF(ISNUMBER(I{fr}),SUM(I{fr}:N{fr}),"n/a")', tot, font=F(size=8), align=CENTER)
            else:
                wf.cell(row=fr, column=15).value = 'n/a'
            wf.row_dimensions[fr].height = max(30, 11 * (len(B.txt(outs)) // 38 + 1))
            fr += 1
        if fr > HR + 1:
            wf.conditional_formatting.add(f'O{HR + 1}:O{fr - 1}', CellIsRule(operator='greaterThan', formula=['1000'],
                                          fill=PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid'),
                                          font=Font(name='Arial', bold=True, color='9C0006')))
            wf.conditional_formatting.add(f'O{HR + 1}:O{fr - 1}', CellIsRule(operator='between', formula=['500', '1000'],
                                          fill=PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid'),
                                          font=Font(name='Arial', bold=True, color='9C5700')))
            dvp = DataValidation(type='list', formula1='"a,b,c,d,e"', allow_blank=True)
            dvp.add(f'P{HR + 1}:P{fr - 1}')
            wf.add_data_validation(dvp)
            dvk = DataValidation(type='list', formula1='"B,1,2,3,4"', allow_blank=True)
            dvk.add(f'Q{HR + 1}:Q{fr - 1}')
            wf.add_data_validation(dvk)
        for i, w in enumerate([7, 8, 24, 34, 40, 8, 30, 30, 9, 8, 10, 10, 10, 11, 11, 9, 9, 13, 36], 1):
            wf.column_dimensions[CL(i)].width = w
        wf.freeze_panes = f'D{HR + 1}'
        wf.auto_filter.ref = f'A{HR}:S{max(HR + 1, fr - 1)}'
        B.put(wf, fr + 1, 1, self.t('sf.cf_note'), F(size=8, italic=True), border=False)

    # ----------------------------------------------------------------- Test record
    def sheet_tests(self):
        B, A, T = self.B, self.A, self.t
        wt = B.wb.create_sheet(self.SN['test'])
        B.put(wt, 1, 1, T('t.title'), F(bold=True, size=12), border=False)
        B.put(wt, 2, 1, self.title_block() + '   •   ' + self.source(), F(size=8, bold=True), border=False)
        B.put(wt, 3, 1, T('t.pre'), F(size=8), None, LEFT, border=False)
        wt.merge_cells('A3:M3')
        wt.row_dimensions[3].height = 30
        B.put(wt, 4, 1, T('t.result_hint') + T('input_note'), F(size=8, italic=True), INFILL, LEFT, border=False)
        wt.merge_cells('A4:M4')
        TR = 6
        for i, h in enumerate(self.tr.table('t.cols')[0], 1):
            B.put(wt, TR, i, h, F(bold=True, color='FFFFFF', size=8), HDR_FILL, CENTER)
        wt.row_dimensions[TR].height = 30
        tests = []

        def short_reset(tag):
            if not tag or INS.looks_literal(tag):
                return self.tr.ml(tag or '—'), ''
            srcs, locs = A.reset_sources(tag)
            nice = [self.src_text(s) for s in srcs if A.source_kind(s) not in ('const0', 'const1', 'never')]
            return (self.tr.join('; ', nice) if nice else self.tr.ml(tag)), ', '.join(self.loc(l) for l in locs)

        def restart_text(effs):
            out = []
            routs = [self.EM[e.key]['label'] for e in effs if e.instr]
            if routs:
                out.append(T('t.rs.rout', outs=', '.join(routs)))
            groups = OrderedDict()
            for e in effs:
                if e.instr or e.info:
                    continue
                m = self.EM[e.key]
                if m.get('auto'):
                    groups.setdefault(('AUTO', ', '.join(m['auto'])), []).append(e)
                elif m.get('manual'):
                    groups.setdefault(('MAN', ' ∥ '.join(m['manual'])), []).append(e)
                else:
                    groups.setdefault(('FOLLOW', ''), []).append(e)
            for (kind, rs), es in groups.items():
                names = self.compress(es)
                if kind == 'AUTO':
                    out.append(T('t.rs.auto', outs=names, tags=rs))
                elif kind == 'MAN':
                    rr = ' ∥ '.join(f"{t} [{short_reset(t)[1]}]" if short_reset(t)[1] else t for t in rs.split(' ∥ '))
                    out.append(T('t.rs.manual', n=len(es), outs=names, resets=rr))
                else:
                    out.append(T('t.rs.follow', outs=names))
            if groups:
                out.append(T('t.rs.where'))
            return self.tr.join(' ', out) if out else self.tr.ml('')

        for tag in self.devs:
            d, dev = self.DM[tag], A.devices[tag]
            hx, xc = self.hits(tag), self.hits(tag, ('XC',))
            xb_bits = sorted({b for e in self.safe for c in [self.cells.get((tag, e.key))] if c and c.raw == 'XB' for b in c.bypass})
            short = d['desc'].map(lambda s: s.split('  [')[0])
            two_ch = (dev.kind == 'instr' and dev.ch_a and dev.ch_b and dev.phys_a != dev.phys_b and not dev.const_channels
                      and not dev.std_channels)
            act = self.tr.ml('1) ') + T('act.' + d['type']) + ': ' + short + '.'
            if xb_bits:
                act = act + ' ' + T('t.act.xb', bits=', '.join(xb_bits))
            if d['mute']:
                act = act + ' ' + T('t.act.mute')
            if two_ch:
                act = act + ' ' + T('t.act.disc', a=dev.phys_a.replace('Data', ''), b=dev.phys_b.replace('Data', ''))
            sto = any(e.kind in DRIVE_KINDS for e in hx)
            codes = {self.code(tag, e.key) for e in self.safe}
            nconst = [e for e in self.safe if self.code(tag, e.key) == 'N' and not dev.const_channels]
            hx = [e for e in hx if e not in nconst]
            if 'N' in codes and dev.const_channels:
                exp = T('t.exp.n', rung=self.loc(dev.loc))
            elif nconst and not hx:
                exp = self.t_nconst(tag, nconst)
            elif hx:
                exp = T('t.exp.off', outs=self.compress(hx)) + (T('t.exp.sto') if sto else self.tr.ml('')) + T('t.exp.rest')
                if xc:
                    partners = sorted({p for e in xc for p in self.cells[(tag, e.key)].partners})
                    exp = exp + ' ' + T('t.exp.notoff', effs=', '.join(self.EM[e.key]['label'] for e in xc), partners=', '.join(partners))
                if two_ch:
                    exp = exp + ' ' + T('t.exp.disc', instr=dev.instr)
                elif dev.equal_channels and dev.kind == 'instr':
                    exp = exp + ' ' + T('t.exp.single')
                if nconst:
                    exp = exp + ' ' + self.t_nconst(tag, nconst)
            elif xc:
                partners = sorted({p for e in xc for p in self.cells[(tag, e.key)].partners})
                exp = T('t.exp.xc_only', effs=', '.join(self.EM[e.key]['label'] for e in xc), partners=', '.join(partners))
                if two_ch:
                    exp = exp + ' ' + T('t.exp.disc_short', instr=dev.instr)
            else:
                exp = T('t.exp.none')
            if dev.reset_manual is True:
                src, rr = short_reset(dev.reset)
                rst = T('t.rst.manual', tag=dev.reset, rr=rr, src=src)
            elif dev.reset_manual is False:
                rst = T('t.rst.auto') + (T('t.rst.auto_estop') if d['type'] == 'EMERGENCY STOP' else self.tr.ml('')) + '.'
            else:
                rst = T('t.rst.follow')
            rst = rst + ' ' + restart_text(hx)
            notes = []
            if xb_bits:
                notes.append(T('t.note.xb', bits=', '.join(xb_bits)))
            if dev.reset_manual is False:
                notes.append(T('t.note.auto'))
            if dev.std_channels:
                notes.append(T('t.note.std'))
            tests.append((T('t.k.func'), f"{d['id']}\n{tag}", short, act, exp, rst, f"{self.loc(dev.loc)} (+ {self.SN['path']})",
                          self.tr.join(' ', notes) if notes else ''))
        # OR combinations
        seen = set()
        for (dv, ek), c in self.cells.items():
            if c.raw != 'XC' or self.A.effect(ek).info:
                continue
            for p in c.partners:
                k = frozenset((dv, p))
                if k in seen:
                    continue
                seen.add(k)
                effs = [e for e in self.safe if self.cells.get((dv, e.key)) and self.cells[(dv, e.key)].raw == 'XC'
                        and p in self.cells[(dv, e.key)].partners]
                ids = [self.DM[t]['id'] for t in (dv, p)]
                rung = ', '.join(sorted({self.loc(l) for e in effs for t_, l in self.cells[(dv, e.key)].path[1:2]}))
                names = ', '.join(self.EM[e.key]['label'] for e in effs)
                tests.append((T('t.k.or'), f'{ids[0]} + {ids[1]}', f'{dv} + {p}', T('t.or.act', a=dv, b=p),
                              T('t.or.exp', outs=names), T('t.or.rst'), rung, T('t.or.note')))
        # standard bypass
        for tag in self.devs:
            bits = OrderedDict()
            for e in self.safe:
                c = self.cells.get((tag, e.key))
                if c and c.raw == 'XB':
                    for b in c.bypass:
                        bits.setdefault(b, []).append(e)
            for b, effs in bits.items():
                rung = ', '.join(sorted({self.loc(l) for e in effs for t_, l in self.cells[(tag, e.key)].path[1:2]}))
                tests.append((T('t.k.byp'), f"{self.DM[tag]['id']}\n{tag}", self.DM[tag]['desc'].map(lambda s: s.split('  [')[0]),
                              T('t.byp.act', bit=b, src=A.std_trace(b) or 'std'), T('t.byp.exp', outs=self.compress(effs)), '—', rung,
                              T('t.byp.note')))
        # EDM
        for e in self.safe:
            if not e.instr:
                continue
            m = self.EM[e.key]
            fb = m['fb'].replace('Data', '')
            act = T('t.edm.act_neg', fb=fb) if str(m['fbtype']).upper() in ('NEGATIVE', '0') else T('t.edm.act_pos', fb=fb)
            src, rr = short_reset(m['reset_tag'])
            tests.append((T('t.k.edm'), m['label'], m['desc'] + f" – feedback {fb} ({m['fbtype']})", act,
                          T('t.edm.exp', tag=e.instr, outs=m['addr']), T('t.edm.rst', reset=m['reset_tag'], rr=rr, src=src),
                          self.loc(A.lm.outputs[e.instr].loc), ''))
        # communication loss per adapter
        adapters = OrderedDict()
        for tag in self.devs:
            dev = A.devices[tag]
            mod = A.module_of(dev.phys_a) if dev.phys_a else None
            if mod:
                ad = self.p.top_adapter(mod)
                adapters.setdefault(ad.name, {'mod': ad, 'devs': [], 'crtl': set()})
                adapters[ad.name]['devs'].append(tag)
                con = mod.connection('input')
                if con and con.reaction_time_limit_ms:
                    adapters[ad.name]['crtl'].add(con.reaction_time_limit_ms)
        for name, g in adapters.items():
            devs = g['devs']
            off = []
            for e in self.safe:
                on_adapter = any((A.module_of(o) and self.p.top_adapter(A.module_of(o)).name == name) for o in e.outputs)
                if on_adapter or not A.ev.value(e.var, dev_off=devs, free=e.baseline):
                    off.append(e)
            ad = g['mod']
            tests.append((T('t.k.comm'), name, T('t.comm.what', name=name, cat=ad.catalog, addr=ad.address or '—'),
                          T('t.comm.act', name=name), T('t.comm.exp', ids=', '.join(self.DM[t]['id'] for t in devs), outs=self.compress(off)),
                          T('t.comm.rst'), T('t.comm.ref'), T('t.comm.note', crtl=', '.join(f'{_g(x)} ms' for x in sorted(g['crtl'])) or 'n/a')))
        drives = [e for e in self.safe if e.kind in DRIVE_KINDS]
        if drives:
            routines = ', '.join(OrderedDict.fromkeys((A.effect_rung(e) or _Nil).routine for e in drives))
            tests.append((T('t.k.comm'), T('t.drv.id'), T('t.drv.what', n=len(drives)), T('t.drv.act'), T('t.drv.exp'), T('t.drv.rst'),
                          T('t.drv.ref', routines=routines), ', '.join(self.EM[e.key]['label'] for e in drives)))
        for n, t in enumerate(tests, 1):
            rr_ = TR + n
            vals = [f'T-{n:03d}'] + list(t) + [None, None, None, None]
            for i, v in enumerate(vals, 1):
                cell = B.put(wt, rr_, i, v, F(size=8, bold=(i in (1, 2))), None, LEFT_TOP)
                if i >= 10:
                    cell.fill = INFILL
                if i == 9 and v and re.search(r'ALARM|\[', B.txt(v)):
                    cell.font = F(size=8, bold=True, color='C00000')
            longest = max(len(B.txt(v)) for v in vals[4:7])
            wt.row_dimensions[rr_].height = min(300, max(26, 10.5 * (longest // 72 + 1)))
        last_t = TR + len(tests)
        self.n_tests = len(tests)
        if tests:
            dvr = DataValidation(type='list', formula1='"OK,NOK,N/A"', allow_blank=True, showErrorMessage=True, error='OK / NOK / N/A')
            dvr.add(f'J{TR + 1}:J{last_t}')
            wt.add_data_validation(dvr)
            for code, (bg, fg) in RESULT_STYLE.items():
                wt.conditional_formatting.add(f'J{TR + 1}:J{last_t}', CellIsRule(operator='equal', formula=[f'"{code}"'],
                                              fill=PatternFill(start_color=bg, end_color=bg, fill_type='solid'),
                                              font=Font(name='Arial', bold=True, color=fg)))
        for i, w in enumerate([7, 14, 22, 28, 48, 52, 50, 18, 26, 8, 10, 12, 22], 1):
            wt.column_dimensions[CL(i)].width = w
        wt.freeze_panes = f'D{TR + 1}'
        wt.auto_filter.ref = f'A{TR}:M{max(TR + 1, last_t)}'
        s_r = last_t + 2
        self.test_range = (TR + 1, last_t)
        B.put(wt, s_r, 1, T('t.summary'), F(bold=True), border=False)
        n = len(tests)
        for k, (lab, formula, val) in enumerate([(T('t.sum.count'), f'=COUNTA(A{TR + 1}:A{last_t})', n),
                                                 (self.tr.ml('OK'), f'=COUNTIF(J{TR + 1}:J{last_t},"OK")', 0),
                                                 (self.tr.ml('NOK'), f'=COUNTIF(J{TR + 1}:J{last_t},"NOK")', 0),
                                                 (self.tr.ml('N/A'), f'=COUNTIF(J{TR + 1}:J{last_t},"N/A")', 0),
                                                 (T('t.sum.open'), f'=COUNTA(A{TR + 1}:A{last_t})-COUNTA(J{TR + 1}:J{last_t})', n)], 1):
            B.put(wt, s_r + k, 2, lab, F(bold=True, size=8), TYPE_FILL, LEFT)
            B.put_number_formula(wt, s_r + k, 3, formula, val, font=F(size=8, bold=True), align=CENTER)
        B.put(wt, s_r + 7, 2, T('t.approve'), F(bold=True, size=8), border=False)
        for k, key in enumerate(('t.done_by', 't.checked_by', 't.approved_by')):
            B.put(wt, s_r + 8 + k, 2, T(key), F(size=8), TYPE_FILL, LEFT)
            B.put(wt, s_r + 8 + k, 3, None, F(size=8), INFILL, LEFT)
            wt.merge_cells(start_row=s_r + 8 + k, start_column=3, end_row=s_r + 8 + k, end_column=6)

    # ----------------------------------------------------------------- Start
    def sheet_start(self):
        B, T = self.B, self.t
        ws = self.ws_start
        Fs = lambda **k: Font(name='Arial', size=k.pop('size', 10), **k)
        B.put(ws, 1, 2, self.name(), Fs(bold=True, size=16), border=False)
        B.put(ws, 2, 2, T('s.subtitle', proc=self.p.processor or '—', ver=self.p.software_version or '—',
                          period=_g(self.A.task_period), src=self.p.source_name or '—'), Fs(size=10, italic=True), border=False)
        B.put(ws, 3, 2, T('s.author', author=self.cfg.author or '—', rev=self.cfg.revision, date=self.cfg.date), Fs(bold=True), border=False)
        for r_ in (1, 2, 3):
            ws.merge_cells(start_row=r_, start_column=2, end_row=r_, end_column=6)
        r = 5
        if self.multi:
            lab = ws.cell(row=5, column=2, value=' / '.join(self.tr.s(l, 's.lang_label') for l in self.B.langs))
            lab.font = Fs(bold=True, size=12)
            lab.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
            sw = ws.cell(row=5, column=3, value=self.B.L0.upper())
            sw.font = Fs(bold=True, size=16, color='1F3864')
            sw.fill = INFILL
            med = Side('medium')
            sw.border = Border(left=med, right=med, top=med, bottom=med)
            sw.alignment = Alignment(horizontal='center', vertical='center')
            codes = ','.join(l.upper() for l in self.B.langs)
            dv = DataValidation(type='list', formula1=f'"{codes}"', allow_blank=False, showErrorMessage=True,
                                errorTitle='Language', error=codes.replace(',', ' / '), showInputMessage=True,
                                promptTitle='Language', prompt=', '.join(f'{l.upper()} = {LANG_NAMES[l]}' for l in self.B.langs))
            dv.add('C5')
            ws.add_data_validation(dv)
            lc = ws.cell(row=5, column=8, value=B.lcol_formula())
            lc.font = Fs(color='FFFFFF', size=8)
            B.cache[(ws.title, 'H5')] = self.B.langs.index(self.B.L0) + 1
            ws.column_dimensions['H'].hidden = True
            B.put(ws, 5, 4, T('s.switch_hint'), Fs(size=10, bold=True, color='C00000'), border=False)
            ws.merge_cells('D5:F5')
            B.put(ws, 6, 2, T('s.dict_note', dict=self.SN['dict']), Fs(size=9, italic=True), border=False)
            ws.merge_cells('B6:F6')
            ws.row_dimensions[6].height = 36
            r = 8
        B.put(ws, r, 2, T('s.sheet'), Fs(bold=True, color='FFFFFF'), HDR_FILL)
        B.put(ws, r, 3, T('s.contents'), Fs(bold=True, color='FFFFFF'), HDR_FILL)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        r += 1
        for name, txt in self.sheet_list():
            c = ws.cell(row=r, column=2, value=name)
            c.font = Fs(bold=True, color='0563C1', underline='single', size=9)
            c.border = BOX
            c.fill = TYPE_FILL
            c.alignment = Alignment(vertical='center', wrap_text=True)
            c.hyperlink = f"#'{name}'!A1"
            B.put(ws, r, 3, txt, Fs(size=9))
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
            ws.row_dimensions[r].height = 30
            r += 1
        r += 1
        B.put(ws, r, 2, T('s.summary'), Fs(bold=True, color='FFFFFF'), HDR_FILL)
        B.put(ws, r, 3, T('s.value'), Fs(bold=True, color='FFFFFF'), HDR_FILL)
        r += 1
        mf, ml_, mlast = getattr(self, 'motor_range', (3, 3, 4))
        cnt = lambda code: f'COUNTIF(MOTOR!$D${mf}:${CL(mlast)}${ml_},"{code}")'
        motor_keys = [e.key for e in self.effects if not self.EM[e.key]['valve']]
        counts = {c: sum(1 for d in self.devs for k in motor_keys if self.code(d, k) == c) for c in ('X', 'X1', 'XB', 'XC', 'XM', 'N')}
        n_info = sum(1 for e in self.effects if e.info)
        t0, t1 = getattr(self, 'test_range', (7, 7))
        summary = [
            (T('s.sum.devices'), len(self.devs), None),
            (T('s.sum.effects'), f'{len(self.safe)} + {n_info}' if n_info else len(self.safe), None),
            (T('s.sum.cells'), '=' + '&" / "&'.join(cnt(c) for c in counts), ' / '.join(str(v) for v in counts.values())),
            (T('s.sum.alarm'), self.n_alarm, None),
            (T('s.sum.dead'), self.n_dead, None),
            (T('s.sum.verify'), self.n_verify, None),
            (T('s.sum.tests'), f"=COUNTA('{self.SN['test']}'!A{t0}:A{t1})", self.n_tests),
        ]
        for lab, val, cached in summary:
            B.put(ws, r, 2, lab, Fs(size=9))
            if cached is not None:
                B.put_number_formula(ws, r, 3, val, cached, font=Fs(bold=True, size=10), align=CENTER)
            else:
                B.put(ws, r, 3, val, Fs(bold=True, size=10), align=CENTER)
            r += 1
        for col, w in zip('ABCDEF', [3, 46, 22, 40, 30, 30]):
            ws.column_dimensions[col].width = w
        ws.row_dimensions[1].height = 28
        ws.row_dimensions[5].height = 32
        ws.sheet_view.showGridLines = False

    # ----------------------------------------------------------------- finish
    def finish(self, path):
        wb, B = self.B.wb, self.B
        author = self.cfg.author or '—'
        wb.properties.creator = author
        wb.properties.lastModifiedBy = author
        wb.properties.title = B.txt(self.name())
        wb.properties.subject = B.txt(self.t('doc.subject'))
        wb.properties.description = f'{B.txt(self.source())}. Rev. {self.cfg.revision} ({self.cfg.date}). safety-matrix-generator {__version__}'
        wb.properties.keywords = 'Safety Matrix; Cause & Effect; GuardLogix; ISO 13849; IEC 62061'
        wb.properties.language = EXCEL_LANG.get(B.L0, 'en-GB')
        wb.calculation.fullCalcOnLoad = True
        for sh in wb.worksheets:
            sh.oddFooter.left.text = f"{author} – {B.txt(self.name())[:60]}, rev. {self.cfg.revision} ({self.cfg.date})"
            sh.oddFooter.left.size = 8
            sh.oddFooter.right.text = '&P / &N  –  &A'
            sh.oddFooter.right.size = 8
            sh.page_setup.orientation = 'landscape'
            sh.page_setup.paperSize = sh.PAPERSIZE_A3
            sh.page_setup.fitToWidth = 1
            sh.page_setup.fitToHeight = 0
            sh.sheet_properties.pageSetUpPr.fitToPage = True
            sh.print_options.gridLines = False
            sh.sheet_view.tabSelected = sh is wb.worksheets[0]
        wb.active = 0
        wb.save(path)
        inject_cached_values(path, [s.title for s in wb.worksheets], B.cache)


class _Nil:
    routine = ''


def _g(v):
    if v is None:
        return 'n/a'
    if isinstance(v, float):
        return f'{v:g}'
    return str(v)


def _runs(nums):
    nums = sorted(set(nums))
    out = []
    for n in nums:
        if out and n == out[-1][1] + 1:
            out[-1][1] = n
        else:
            out.append([n, n])
    return out


def _ranges(rows, cols):
    return ' '.join(f'{CL(c0)}{r0}:{CL(c1)}{r1}' for r0, r1 in _runs(rows) for c0, c1 in _runs(cols))


def _compress_outputs(outs: List[str]) -> str:
    """['M:6:O.Pt02Data', 'M:6:O.Pt03Data'] -> 'M:6:O.Pt02/Pt03'."""
    if not outs:
        return ''
    parts = [io_parts(o) for o in outs]
    if all(parts) and len({(p[0], p[1]) for p in parts}) == 1:
        return f"{parts[0][0]}:{parts[0][1]}." + '/'.join(p[2].replace('Data', '') for p in parts)
    return ', '.join(outs)
