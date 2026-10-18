# =============================================================================
# Project : Safety Matrix Generator (GuardLogix SafetyTask -> Cause & Effect)
# File    : safety_matrix/excel/base.py
# Author  : Robert Pietroń
# Revision: 1.0.0 (2026-10-18)
# Changes :
#   1.0.0 (2026-10-18) - First universal release: styles, multi-language cell
#                        writer (dictionary sheet + INDEX formulas driven by the
#                        language switch) and injection of cached formula
#                        results so every viewer shows the default language.
# =============================================================================
"""Workbook primitives shared by all sheets."""
from __future__ import annotations

import re
import shutil
import tempfile
import zipfile
from typing import Dict, List, Tuple
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

from ..i18n import ML, Tr

# ------------------------------------------------------------------ styles
def F(**k):
    return Font(name='Arial', size=k.pop('size', 9), **k)


thin = Side(style='thin', color='808080')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill('solid', fgColor='1F3864')
ZONE_FILL = PatternFill('solid', fgColor='D9E1F2')
AREA_FILL = PatternFill('solid', fgColor='FCE4D6')
TYPE_FILL = PatternFill('solid', fgColor='EDEDED')
META_FILL = PatternFill('solid', fgColor='F2F2F2')
INPUT_FILL_HEX = 'FFF2CC'
INFILL = PatternFill('solid', fgColor=INPUT_FILL_HEX)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
VERT = Alignment(horizontal='center', vertical='bottom', text_rotation=90, wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
LEFT_TOP = Alignment(horizontal='left', vertical='top', wrap_text=True)

# code, background, text colour, readable colour on white
CODE_STYLE = {
    'X': ('C00000', 'FFFFFF', 'C00000'), 'X1': ('F4B084', '000000', 'C55A11'), 'XB': ('FF9900', '000000', 'E26B0A'),
    'XC': ('FFE699', '000000', 'BF8F00'), 'XM': ('D9D2E9', '000000', '7030A0'), 'N': ('A6A6A6', '000000', '7F7F7F'),
    'S': ('9BC2E6', '000000', '2F75B5'), 'SB': ('FF9900', '000000', 'E26B0A'), 'SC': ('FFE699', '000000', 'BF8F00'),
    'SM': ('D9D2E9', '000000', '7030A0'), '': ('FFFFFF', '000000', '000000'),
}
RESULT_STYLE = {'OK': ('C6EFCE', '006100'), 'NOK': ('FFC7CE', '9C0006'), 'N/A': ('EDEDED', '000000')}


def IFont(sz=9, b=False, i=False, color='000000'):
    return InlineFont(rFont='Arial', sz=sz, b=b, i=i, color=color)


class Book:
    """openpyxl workbook with multi-language text cells."""

    def __init__(self, tr: Tr, default_lang: str, author: str):
        self.tr = tr
        self.langs = tr.langs
        self.L0 = default_lang if default_lang in self.langs else self.langs[0]
        self.multi = len(self.langs) > 1
        self.author = author or '—'
        self.wb = Workbook()
        self.entries: List[Tuple[str, ...]] = []
        self.index: Dict[Tuple[str, ...], int] = {}
        self.first_use: Dict[int, str] = {}
        self.cache: Dict[Tuple[str, str], object] = {}   # (sheet, coord) -> cached formula result

    # ----------------------------------------------------------------- text handling
    def txt(self, v, lang=None) -> str:
        if isinstance(v, ML):
            return v.get(lang or self.L0, v.first())
        return '' if v is None else str(v)

    def _key(self, v: ML) -> int:
        k = tuple(v.get(l, '') for l in self.langs)
        if k not in self.index:
            self.entries.append(k)
            self.index[k] = len(self.entries)
        return self.index[k]

    @property
    def fallback_col(self) -> int:
        return self.langs.index('en') + 1 if 'en' in self.langs else 1

    def formula(self, n: int) -> str:
        fb = self.fallback_col
        return f'=IF(INDEX(TXT,{n},LCOL)="",INDEX(TXT,{n},{fb}),INDEX(TXT,{n},LCOL))'

    def set_value(self, cell, v):
        ws = cell.parent
        if isinstance(v, ML):
            if not self.multi or v.same():
                cell.value = self.txt(v)
            else:
                n = self._key(v)
                self.first_use.setdefault(n, f'{ws.title}!{cell.coordinate}')
                cell.value = self.formula(n)
                self.cache[(ws.title, cell.coordinate)] = self.txt(v)
        else:
            cell.value = v

    def put(self, ws, r, c, v, font=None, fill=None, align=None, border=True):
        cell = ws.cell(row=r, column=c)
        self.set_value(cell, v)
        cell.font = font or F()
        if fill:
            cell.fill = fill
        if align:
            cell.alignment = align
        if border:
            cell.border = BOX
        return cell

    def put_number_formula(self, ws, r, c, formula: str, value=None, **kw):
        cell = self.put(ws, r, c, formula, **kw)
        if value is not None:
            self.cache[(ws.title, cell.coordinate)] = value
        return cell

    def comment(self, cell, v, width=420, height=140):
        txt = self.txt(v)
        if self.multi and isinstance(v, ML) and 'en' in v and self.L0 != 'en' and v['en'] != txt:
            txt = txt + '\n—\n' + v['en']
        cm = Comment(txt, self.author)
        cm.width, cm.height = width, height
        cell.comment = cm

    def rich(self, ws, rng: str, parts: List[Tuple[object, InlineFont]]):
        """Rich text block (single language) or a plain multi-language formula (switchable)."""
        ws.merge_cells(rng)
        first = ws[rng.split(':')[0]]
        if self.multi:
            self.set_value(first, self.tr.of(lambda l: ''.join(self.txt(t, l) for t, _ in parts)))
        else:
            first.value = CellRichText([TextBlock(f, self.txt(t)) for t, f in parts])
        first.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

    # ----------------------------------------------------------------- finishing
    def sheet_name(self, key: str) -> str:
        """Sheet name in the default language, plus English when several languages are written."""
        a = self.tr.s(self.L0, key)
        if self.multi and self.L0 != 'en' and 'en' in self.langs:
            b = self.tr.s('en', key)
            if b != a and len(f'{a} | {b}') <= 31:
                return f'{a} | {b}'
        return a[:31]

    def write_dictionary(self, sheet_title: str, start_sheet: str, switch_cell='C5', lcol_cell='H5'):
        from openpyxl.workbook.defined_name import DefinedName
        wd = self.wb.create_sheet(sheet_title)
        hdr = ['No.'] + [l.upper() for l in self.langs] + ['First use']
        for i, h in enumerate(hdr, 1):
            self.put(wd, 1, i, h, F(bold=True, color='FFFFFF', size=9), HDR_FILL, CENTER)
        for n, vals in enumerate(self.entries, 1):
            row = [n, *vals, self.first_use.get(n, '')]
            for i, v in enumerate(row, 1):
                c = wd.cell(row=n + 1, column=i, value=v)
                c.font = F(size=8)
                c.alignment = LEFT_TOP
                c.border = BOX
                if 1 < i <= len(self.langs) + 1 and self.langs[i - 2] not in ('pl', 'en'):
                    c.fill = INFILL
        for i, w in enumerate([7] + [52] * len(self.langs) + [24], 1):
            wd.column_dimensions[CL(i)].width = w
        wd.freeze_panes = 'B2'
        wd.auto_filter.ref = f'A1:{CL(len(hdr))}{len(self.entries) + 1}'
        last = CL(len(self.langs) + 1)
        self.wb.defined_names['LANG'] = DefinedName('LANG', attr_text=f"'{start_sheet}'!${switch_cell[0]}${switch_cell[1:]}")
        self.wb.defined_names['LCOL'] = DefinedName('LCOL', attr_text=f"'{start_sheet}'!${lcol_cell[0]}${lcol_cell[1:]}")
        self.wb.defined_names['TXT'] = DefinedName('TXT', attr_text=f"'{sheet_title}'!$B$2:${last}${max(2, len(self.entries) + 1)}")
        return wd

    def lcol_formula(self) -> str:
        expr = '1'
        for i in range(len(self.langs) - 1, 0, -1):
            expr = f'IF(LANG="{self.langs[i].upper()}",{i + 1},{expr})'
        return '=' + expr


def inject_cached_values(path: str, sheet_titles: List[str], cache: Dict[Tuple[str, str], object]):
    """Write cached results (<v>) for formula cells so viewers that do not recalculate show the text."""
    if not cache:
        return
    by_sheet: Dict[str, Dict[str, object]] = {}
    for (sh, coord), v in cache.items():
        by_sheet.setdefault(sh, {})[coord] = v
    tmp = tempfile.mktemp(suffix='.xlsx')
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            m = re.match(r'xl/worksheets/sheet(\d+)\.xml$', item.filename)
            if m:
                idx = int(m.group(1)) - 1
                title = sheet_titles[idx] if idx < len(sheet_titles) else None
                cells = by_sheet.get(title)
                if cells:
                    xml = data.decode('utf-8')

                    def repl(mo):
                        coord = mo.group(1)
                        if coord not in cells:
                            return mo.group(0)
                        attrs, f = mo.group(2), mo.group(3)
                        v = cells[coord]
                        attrs = re.sub(r'\s+t="[^"]*"', '', attrs)
                        if isinstance(v, (int, float)) and not isinstance(v, bool):
                            return f'<c r="{coord}"{attrs}><f>{f}</f><v>{v}</v></c>'
                        return f'<c r="{coord}"{attrs} t="str"><f>{f}</f><v>{escape(str(v))}</v></c>'

                    xml = re.sub(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f>(?:<v>.*?</v>|<v\s*/>)?</c>', repl, xml, flags=re.S)
                    data = xml.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
