# =============================================================================
# Project : Safety Matrix Generator (GuardLogix SafetyTask -> Cause & Effect)
# File    : tests/test_rules_v11.py
# Author  : Robert Pietroń
# Revision: 1.0.0 (2026-10-21)
# Changes :
#   1.0.0 (2026-10-21) - Tests of the 1.1.0 analysis and audit rules on the
#                        synthetic ScanCell export (copied channels, scan lag,
#                        constant bypass, timer presets from tag values, ROUT
#                        via copies, standard writes via AOI / ST / FBD, ...).
# =============================================================================
import functools
import os
import tempfile
import unittest

from openpyxl import load_workbook

from safety_matrix.analysis import Analysis
from safety_matrix.audit import ALARM, DEAD, VERIFY, run_audit
from safety_matrix.model import Tag, split_operand
from safety_matrix.parsers import load_project
from safety_matrix.parsers.l5x import parse_l5x
from safety_matrix.pipeline import run

from .helpers import ROOT

SCAN = os.path.join(ROOT, 'tests', 'data', 'ScanCell.L5K')


@functools.lru_cache(maxsize=None)
def scan():
    A = Analysis(load_project(SCAN))
    return A, run_audit(A)


def keys(F, sev=None):
    return [f.key for f in F if sev is None or f.severity == sev]


def find(F, key):
    return [f for f in F if f.key == key]


class TestModelHelpers(unittest.TestCase):
    def test_split_operand_program_scope(self):
        self.assertEqual(split_operand('SafetyProgram\\IBN.0'), ('SafetyProgram\\IBN', '.0'))
        self.assertEqual(split_operand('Local:1:I.Pt00Data'), ('Local:1:I', '.Pt00Data'))
        self.assertEqual(split_operand('Arr[2].5'), ('Arr', '[2].5'))

    def test_timer_preset(self):
        self.assertEqual(Tag('T', 'TIMER', value=[0, 1500, 0]).timer_preset(''), 1500.0)
        self.assertEqual(Tag('T', 'TIMER[3]', value=[[0, 0, 0], [0, 2500, 0], [0, 7, 0]]).timer_preset('[1]'), 2500.0)
        self.assertIsNone(Tag('T', 'TIMER', value=None).timer_preset(''))


class TestChannelsAndLag(unittest.TestCase):
    def test_copied_channels_resolved(self):
        A, _ = scan()
        d = A.devices['EStop_A']
        self.assertEqual((d.ch_a, d.src_a, d.phys_b), ('Ch_A', 'Local:1:I.Pt00Data', 'Local:1:I.Pt01Data'))
        self.assertEqual(d.channel_lag, 1)                     # copies run after the ESTOP instruction
        m, con = A.input_connection(d)
        self.assertEqual((m.name, con.reaction_time_limit_ms), ('SafeIn_1', 40.064))
        self.assertEqual(A.io_comment(d.phys_a), '-S1 E-stop ch.A')

    def test_scan_lag_in_cells(self):
        A, F = scan()
        self.assertEqual(A.cells[('Door_1', 'Local:2:O.Pt07Data')].lag, 1)      # Late_OK written later
        self.assertEqual(A.cells[('EStop_A', 'Local:2:O.Pt02Data')].lag, 1)     # channel copy lag
        self.assertEqual(A.cells[('Door_1', 'K_Main')].lag, 0)
        lag = find(F, 'scan_lag')
        self.assertEqual([f.element for f in lag], ['Late_OK'])
        late = find(F, 'channel_copy_late')
        self.assertEqual(len(late), 1)
        self.assertIn('Ch_A ← Local:1:I.Pt00Data', late[0].params['devices'])
        crtl = find(F, 'crtl_inputs')[0]
        self.assertEqual((crtl.params['lags'], crtl.params['total']), (1, '160.128'))


class TestCells(unittest.TestCase):
    def test_constant_bypass_is_N_and_alarm(self):
        A, F = scan()
        c = A.cells[('Door_1', 'Local:2:O.Pt03Data')]
        self.assertEqual((c.code, c.bypass), ('N', ['SafetyProgram\\IBN.0']))
        cb = find(F, 'const_branch')
        self.assertEqual(len(cb), 1)
        self.assertEqual((cb[0].severity, cb[0].params['devices']), (ALARM, 'Door_1'))
        unused = ' '.join(f.element for f in find(F, 'unused_tags'))
        self.assertNotIn('IBN', unused)

    def test_rout_through_copy_grouped(self):
        A, F = scan()
        k = A.effect('K_Main')
        self.assertEqual((k.kind, k.outputs), ('ROUT', ['Local:2:O.Pt00Data', 'Local:2:O.Pt01Data']))
        self.assertEqual({A.code('EStop_A', 'K_Main'), A.code('Door_1', 'K_Main')}, {'X'})
        self.assertEqual(find(F, 'rout_extra_logic')[0].params['devices'], 'Door_1')
        self.assertEqual(find(F, 'edm_same_channel')[0].params['ch'], 'Local:1:I.Pt04Data')

    def test_timer_presets_from_tags(self):
        A, F = scan()
        self.assertEqual(A.lm.timers['Timers[1]'].preset, 2500.0)
        self.assertTrue(A.lm.timers['Timers[1]'].from_tag)
        self.assertEqual(A.cells[('EStop_A', 'Drive_1:SO.SafeTorqueOff')].delay_ms, 2500.0)
        self.assertIsNone(A.cells[('Door_1', 'Local:2:O.Pt04Data')].delay_ms)
        self.assertEqual(find(F, 'stop_cat1')[0].params['effects'], 'Drive_1:SO.SafeTorqueOff')
        other = find(F, 'timer_other')
        self.assertEqual((other[0].params['preset'], other[0].params['cond']), ('4000', 'NOT S_Fbd'))


class TestAuditRules(unittest.TestCase):
    def test_no_std_parallel_for_series_contact(self):
        _, F = scan()
        self.assertNotIn('std_parallel', keys(F))
        self.assertNotIn('std_bypass', keys(F))

    def test_std_writers(self):
        A, F = scan()
        self.assertEqual([f.element for f in find(F, 'std_not_written')], ['S_None'])
        self.assertEqual(A.std_writers('Std_Permit')[0][3], 'MOTORCTL')          # AOI output parameter
        self.assertEqual(A.std_writers('Std_St')[0][3], 'ST')
        self.assertEqual(A.std_writers('Std_Fbd')[0][3], 'FBD')
        self.assertEqual(A.std_writers('Std_None'), [])

    def test_map_target_written(self):
        _, F = scan()
        mw = find(F, 'map_written')
        self.assertEqual([(f.severity, f.element) for f in mw], [(VERIFY, 'S_Coil')])
        self.assertNotIn('map_read_before_write', keys(F))

    def test_aux_outputs(self):
        _, F = scan()
        self.assertEqual([f.element for f in find(F, 'aux_always_on')], ['Local:2:O.Pt05Data'])
        inv = find(F, 'aux_inverse')
        self.assertEqual((inv[0].element, inv[0].params['devices']), ('Local:2:O.Pt06Data', 'EStop_A'))
        self.assertNotIn('aux_no_device', keys(F))

    def test_afi_without_outputs(self):
        _, F = scan()
        self.assertEqual(keys(F, ALARM).count('afi_noop'), 1)
        self.assertNotIn('afi', keys(F))

    def test_drive_peripheral_not_a_drive(self):
        _, F = scan()
        hw = find(F, 'drives_hardwired')
        self.assertEqual([(f.severity, f.element) for f in hw], [(DEAD, 'Drive_HW (PowerFlex 525-EENET)')])


class TestWorkbookV11(unittest.TestCase):
    def test_workbook_texts(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, 'scan.xlsx')
            run(SCAN, out, None, ['en'], 'en')
            wb = load_workbook(out, data_only=True)

            def row_with(sheet, text):
                for row in wb[sheet].iter_rows(values_only=True):
                    if any(v == text for v in row):
                        return [str(v) for v in row if v not in (None, '')]
                return []

            outs = ' | '.join(row_with('Outputs', 'K_Main'))
            self.assertIn('restart only with Reset_Edge', outs)
            devs = ' | '.join(row_with('Devices', 'EStop_A'))
            self.assertIn('Ch_A ← Local:1:I.Pt00Data', devs)
            path = [r for r in wb['C&E paths'].iter_rows(values_only=True) if r[0] == 'Door_1' and r[2] == 'N']
            self.assertEqual(len(path), 1)
            self.assertIn('SafetyProgram\\IBN.0 = 1 (const)', path[0][4])


class TestL5XNonLadderSource(unittest.TestCase):
    def test_st_and_fbd_source(self):
        xml = '''<?xml version="1.0" encoding="UTF-8"?>
<RSLogix5000Content SoftwareRevision="33.01"><Controller Name="X" ProcessorType="5069-L320ERMS2">
<Programs><Program Name="P" MainRoutineName="R"><Tags/><Routines>
<Routine Name="S" Type="ST"><STContent><Line Number="0"><![CDATA[A_Tag := 1;]]></Line></STContent></Routine>
<Routine Name="F" Type="FBD"><FBDContent><Sheet Number="1"><OREF ID="1" Operand="B_Tag"/></Sheet></FBDContent></Routine>
</Routines></Program></Programs></Controller></RSLogix5000Content>'''
        p = parse_l5x(xml)
        self.assertEqual(p.programs['P'].routines['S'].source, ['ST:A_Tag := 1;'])
        self.assertEqual(p.programs['P'].routines['F'].source, ['W:B_Tag'])


if __name__ == '__main__':
    unittest.main()
