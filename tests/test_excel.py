# =============================================================================
# Project : Safety Matrix Generator (GuardLogix SafetyTask -> Cause & Effect)
# File    : tests/test_excel.py
# Author  : Robert Pietroń
# Revision: 1.0.0 (2026-10-18)
# Changes :
#   1.0.0 (2026-10-18) - Workbook generation in four languages and in a single
#                        language; language switch formulas and cached values.
# =============================================================================
import json
import os
import tempfile
import unittest
import zipfile

from openpyxl import load_workbook

from safety_matrix.pipeline import run

from .helpers import DEMO_CFG, DEMO_L5K


class TestWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.multi = os.path.join(cls.tmp.name, 'demo_multi.xlsx')
        cls.single = os.path.join(cls.tmp.name, 'demo_en.xlsx')
        cls.json = os.path.join(cls.tmp.name, 'demo.json')
        cls.sum_multi = run(DEMO_L5K, cls.multi, DEMO_CFG, json_path=cls.json)
        cls.sum_single = run(DEMO_L5K, cls.single, None, ['en'], 'en')

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_sheets_multi(self):
        wb = load_workbook(self.multi)
        self.assertEqual(wb.sheetnames, ['Start', 'Group Interlocks', 'MOTOR', 'VALVE', 'Legenda | Legend',
                                         'Funkcje PL | Safety functions', 'Protokół testów | Test record', 'Urządzenia | Devices',
                                         'Wyjścia | Outputs', 'Ścieżki C&E | C&E paths', 'Raport | Validation report',
                                         'Słownik | Dictionary'])
        self.assertIn('TXT', wb.defined_names)
        self.assertEqual(wb['Start']['C5'].value, 'PL')
        self.assertEqual(wb.properties.creator, 'Robert Pietroń')
        d = wb['Słownik | Dictionary']
        self.assertEqual([d.cell(1, c).value for c in range(2, 6)], ['PL', 'EN', 'FR', 'DA'])

    def test_valve_sheet_and_codes(self):
        wb = load_workbook(self.multi)
        wv = wb['VALVE']
        labels = [wv.cell(2, c).value for c in range(4, wv.max_column + 1)]
        self.assertIn('Valve_Press', labels)
        wm = wb['MOTOR']
        row = next(r for r in range(3, wm.max_row + 1) if str(wm.cell(r, 3).value).endswith('Gate_Infeed'))
        col = next(c for c in range(4, wm.max_column + 1) if wm.cell(2, c).value == 'Robot_Drive')
        self.assertEqual(wm.cell(row, col).value, 'X1')
        self.assertIn('Gate_Infeed.O1', wm.cell(row, col).comment.text)

    def test_cached_values_for_formulas(self):
        with zipfile.ZipFile(self.multi) as z:
            xml = z.read('xl/worksheets/sheet1.xml').decode('utf-8')
        self.assertIn('t="str"><f>IF(INDEX(TXT', xml)
        self.assertIn('Macierz przyczynowo-skutkowa', xml)

    def test_single_language(self):
        wb = load_workbook(self.single)
        self.assertNotIn('Dictionary', wb.sheetnames)
        self.assertIn('Validation report', wb.sheetnames)
        ws = wb['Validation report']
        self.assertTrue(str(ws['A1'].value).startswith('VALIDATION REPORT'))
        self.assertFalse(any(str(c.value).startswith('=IF(INDEX') for row in ws.iter_rows() for c in row if c.value))

    def test_summary_and_json(self):
        self.assertEqual(self.sum_multi['devices'], 9)
        self.assertEqual(self.sum_multi['effects'], 6)
        self.assertEqual(self.sum_single['effects'], 5)
        data = json.load(open(self.json, encoding='utf-8'))
        self.assertEqual(data['generator']['author'], 'Robert Pietroń')
        self.assertEqual(len(data['cells']), 22)


if __name__ == '__main__':
    unittest.main()
